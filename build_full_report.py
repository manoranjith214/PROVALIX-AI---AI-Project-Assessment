import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPORT_DIR = r"c:\Users\Manoranjith\Downloads\PROVALIX-AI\project-report"

FILES_ORDER = [
    "01_FRONT_MATTER.md",
    "02_ABSTRACT.md",
    "03_TABLE_OF_CONTENTS.md",
    "04_CHAPTER_1_INTRODUCTION.md",
    "05_CHAPTER_2_LITERATURE_SURVEY.md",
    "06_CHAPTER_3_MATERIALS_AND_METHODOLOGY.md",
    "07_CHAPTER_4_IMPLEMENTATION.md",
    "08_CHAPTER_5_RESULTS_AND_DISCUSSION.md",
    "09_CHAPTER_6_CONCLUSION_AND_FUTURE_WORK.md",
    "10_REFERENCES.md",
    "11_PROJECT_OUTCOME.md",
    "12_TEST_CASES_AND_EVIDENCE.md",
    "13_SCREENSHOT_CHECKLIST.md"
]

def consolidate_markdown():
    combined_content = []
    print("Consolidating Markdown files...")
    for fname in FILES_ORDER:
        fpath = os.path.join(REPORT_DIR, fname)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as f:
                content = f.read()
                combined_content.append(content)
                combined_content.append("\n\n---\n\\pagebreak\n---\n\n")
        else:
            print(f"Warning: {fname} not found!")

    final_text = "".join(combined_content)
    out_file = os.path.join(REPORT_DIR, "COMPLETE_PROJECT_REPORT.md")
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(final_text)
    print(f"Consolidated Markdown saved to: {out_file} ({len(final_text)} characters)")
    return final_text

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_docx():
    print("Generating Academic DOCX Report...")
    doc = Document()

    # Configure A4 Paper Size
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    # Configure Normal Style (Times New Roman, 12pt, 1.5 spacing, Justified)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal_style.paragraph_format.space_after = Pt(6)

    # Process each markdown file cleanly
    for fname in FILES_ORDER:
        fpath = os.path.join(REPORT_DIR, fname)
        if not os.path.exists(fpath):
            continue

        with open(fpath, "r", encoding="utf-8") as f:
            lines = f.readlines()

        i = 0
        in_table = False
        table_rows = []

        while i < len(lines):
            line = lines[i].rstrip()
            
            # Check for table
            if line.startswith("|") and line.endswith("|"):
                table_rows.append(line)
                i += 1
                continue
            elif table_rows:
                # Render accumulated table
                render_docx_table(doc, table_rows)
                table_rows = []

            # Pagebreak
            if "\\pagebreak" in line or line.strip() == "---":
                if "\\pagebreak" in line:
                    doc.add_page_break()
                i += 1
                continue

            # Skip HTML tags
            if line.startswith("<div") or line.startswith("</div") or line.startswith("<table") or line.startswith("</table") or line.startswith("<tr") or line.startswith("</tr") or line.startswith("<td") or line.startswith("</td"):
                i += 1
                continue

            # Skip Mermaid code blocks
            if line.startswith("```mermaid") or line.startswith("```"):
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    i += 1
                i += 1
                continue

            # Heading 1
            if line.startswith("# "):
                text = line[2:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ("CERTIFICATE" in text or "DECLARATION" in text or "PROVALIX" in text or "ABSTRACT" in text or "TABLE OF CONTENTS" in text) else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(15, 23, 42)
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(12)
                p.paragraph_format.keep_with_next = True
                i += 1
                continue

            # Heading 2
            if line.startswith("## "):
                text = line[3:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text)
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(30, 41, 59)
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                i += 1
                continue

            # Heading 3
            if line.startswith("### "):
                text = line[4:].strip()
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(51, 65, 85)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                i += 1
                continue

            # Bullet points
            if line.startswith("- ") or line.startswith("* "):
                text = line[2:].strip()
                p = doc.add_paragraph(style='List Bullet')
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.3
                add_formatted_runs(p, text)
                i += 1
                continue

            # Numbered lists (e.g. 1. )
            num_match = re.match(r"^(\d+)\.\s+(.*)", line)
            if num_match:
                p = doc.add_paragraph(style='List Number')
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.3
                add_formatted_runs(p, num_match.group(2))
                i += 1
                continue

            # Regular paragraph
            if line.strip():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.5
                add_formatted_runs(p, line.strip())

            i += 1

        if table_rows:
            render_docx_table(doc, table_rows)
            table_rows = []

        doc.add_page_break()

    out_docx = os.path.join(REPORT_DIR, "Provalix_AI_Academic_Project_Report.docx")
    doc.save(out_docx)
    print(f"DOCX Report successfully saved to: {out_docx}")

def add_formatted_runs(paragraph, text):
    # Regex parse bold (**text**) and italics (*text*)
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)

def render_docx_table(doc, table_rows):
    if len(table_rows) < 2:
        return

    # Parse rows
    parsed = []
    for r in table_rows:
        cells = [c.strip() for c in r.split("|")[1:-1]]
        if cells and not all(re.match(r"^:?-+:?$", c) for c in cells):
            parsed.append(cells)

    if not parsed:
        return

    num_cols = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row_data in enumerate(parsed):
        row = table.rows[row_idx]
        is_header = (row_idx == 0)
        for col_idx in range(num_cols):
            cell = row.cells[col_idx]
            text = row_data[col_idx] if col_idx < len(row_data) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if is_header:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_background(cell, "1E293B")
            else:
                run.font.color.rgb = RGBColor(15, 23, 42)
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")

            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

if __name__ == "__main__":
    consolidate_markdown()
    generate_docx()
